"""Turn pandoc's raw.docx into an IEEE-conference-styled two-column Word document."""
import copy, re, sys
import docx
from docx.shared import Pt, Inches, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lxml import etree

RAW, OUT, IMGDIR = sys.argv[1], sys.argv[2], sys.argv[3]
d = docx.Document(RAW)
body = d.element.body
W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
M_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
TNR = 'Times New Roman'
COL_TW, TEXT_TW = 5040, 10440           # column / full text width in twips (3.5 in / 7.25 in)

# ------------------------------------------------------------------ styles
for rf in d.styles.element.iter(qn('w:rFonts')):
    for a in list(rf.attrib):
        if a.endswith('Theme'):
            del rf.attrib[a]
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), TNR)
for c in list(d.styles.element.iter(qn('w:color'))):
    c.getparent().remove(c)


def get_style(name):
    sid = name.replace(' ', '')
    for st in d.styles:
        if st.style_id == sid or st.name == name:
            return st
    return None


def style_fmt(name, size=10, bold=None, italic=None, align=None, before=0, after=0,
              first=None, small_caps=None, keep=False, font=TNR):
    st = get_style(name)
    if st is None:
        return
    f = st.font
    f.name = font; f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if small_caps is not None: f.small_caps = small_caps
    if hasattr(st, 'paragraph_format'):
        pf = st.paragraph_format
        if align is not None: pf.alignment = align
        pf.space_before = Pt(before); pf.space_after = Pt(after)
        pf.line_spacing = 1.0
        if first is not None: pf.first_line_indent = first
        pf.keep_with_next = keep


J = WD_ALIGN_PARAGRAPH.JUSTIFY
for n in ('Normal', 'Body Text', 'First Paragraph', 'Compact'):
    style_fmt(n, 10, align=J, first=Inches(0.17))
style_fmt('Heading 1', 10, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.CENTER, before=8, after=4,
          first=Inches(0), small_caps=True, keep=True)
style_fmt('Heading 2', 10, bold=False, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=5, after=3,
          first=Inches(0), keep=True)
for cs in ('Verbatim Char',):
    st = get_style(cs)
    if st is not None: st.font.name = 'Courier New'; st.font.size = Pt(9)
hl = get_style('Hyperlink')
if hl is not None: hl.font.underline = False

# ------------------------------------------------------------------ helpers
def mkp(text='', size=10, align=None, bold=False, italic=False, small_caps=False, after=0, before=0):
    p = OxmlElement('w:p')
    para = docx.text.paragraph.Paragraph(p, d._body)
    pf = para.paragraph_format
    pf.first_line_indent = Inches(0); pf.space_after = Pt(after); pf.space_before = Pt(before)
    if align is not None: pf.alignment = align
    if text:
        r = para.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
        r.font.small_caps = small_caps; r.font.name = TNR
    return p, para


def sect_break(p_el, cols, first=False):
    """make paragraph p_el the last paragraph of a section with `cols` columns."""
    pPr = p_el.find(qn('w:pPr'))
    if pPr is None:
        pPr = OxmlElement('w:pPr'); p_el.insert(0, pPr)
    sp = OxmlElement('w:spacing'); sp.set(qn('w:before'), '0'); sp.set(qn('w:after'), '0')
    sp.set(qn('w:line'), '20'); sp.set(qn('w:lineRule'), 'exact')
    pPr.append(sp)
    pPr.append(make_sectpr(cols, first))


def make_sectpr(cols, first=False):
    s = OxmlElement('w:sectPr')
    t = OxmlElement('w:type'); t.set(qn('w:val'), 'nextPage' if first else 'continuous'); s.append(t)
    pg = OxmlElement('w:pgSz'); pg.set(qn('w:w'), '12240'); pg.set(qn('w:h'), '15840'); s.append(pg)
    m = OxmlElement('w:pgMar')
    for k, v in (('top', 1080), ('right', 900), ('bottom', 1440), ('left', 900), ('header', 720),
                 ('footer', 720), ('gutter', 0)):
        m.set(qn('w:' + k), str(v))
    s.append(m)
    c = OxmlElement('w:cols'); c.set(qn('w:num'), str(cols)); c.set(qn('w:space'), '360'); s.append(c)
    return s


def para_text(p_el):
    return ''.join(t.text or '' for t in p_el.iter(qn('w:t')))


def strip_marker(p_el, marker):
    for t in p_el.iter(qn('w:t')):
        if t.text and marker in t.text:
            t.text = t.text.replace(marker, '', 1).lstrip()
            return


def runs_fmt(p_el, size=None, bold=None, italic=None, small_caps=None):
    para = docx.text.paragraph.Paragraph(p_el, d._body)
    for r in para.runs:
        if size: r.font.size = Pt(size)
        if bold is not None: r.bold = bold
        if italic is not None: r.italic = italic
        if small_caps is not None: r.font.small_caps = small_caps
    return para


def picture_par(fname, width_in):
    p, para = mkp(align=WD_ALIGN_PARAGRAPH.CENTER, after=2, before=4)
    para.add_run().add_picture(f'{IMGDIR}/{fname}', width=Inches(width_in))
    return p


def booktabs(tbl, widths_cm, wide):
    total = TEXT_TW if wide else COL_TW
    ncol = len(tbl.columns)
    if widths_cm and len(widths_cm) == ncol:
        w = [max(float(x), 0.6) for x in widths_cm]
    else:
        w = [1.6] + [1.0] * (ncol - 1)
    tw = [int(total * x / sum(w)) for x in w]
    fsz = 7 if ncol >= 6 and not wide else 8
    t = tbl._tbl
    tblPr = t.tblPr
    for tag in ('w:tblStyle', 'w:tblW', 'w:tblLayout', 'w:tblBorders', 'w:tblLook', 'w:jc'):
        for e in tblPr.findall(qn(tag)):
            tblPr.remove(e)
    tW = OxmlElement('w:tblW'); tW.set(qn('w:w'), str(sum(tw))); tW.set(qn('w:type'), 'dxa'); tblPr.append(tW)
    jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'center'); tblPr.append(jc)
    b = OxmlElement('w:tblBorders')
    for side, sz in (('top', 8), ('bottom', 8)):
        e = OxmlElement('w:' + side); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:color'), '000000')
        b.append(e)
    for side in ('left', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + side); e.set(qn('w:val'), 'nil'); b.append(e)
    tblPr.append(b)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    cm = OxmlElement('w:tblCellMar')
    for side, v in (('left', 30), ('right', 30)):
        e = OxmlElement('w:' + side); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); cm.append(e)
    tblPr.append(cm)
    grid = t.find(qn('w:tblGrid'))
    if grid is not None:
        for gc, wv in zip(grid.findall(qn('w:gridCol')), tw):
            gc.set(qn('w:w'), str(wv))
    for ri, row in enumerate(t.findall(qn('w:tr'))):
        ci = 0
        for tc in row.findall(qn('w:tc')):
            tcPr = tc.find(qn('w:tcPr'))
            if tcPr is None:
                tcPr = OxmlElement('w:tcPr'); tc.insert(0, tcPr)
            span = tcPr.find(qn('w:gridSpan'))
            n = int(span.get(qn('w:val'))) if span is not None else 1
            for e in tcPr.findall(qn('w:tcW')): tcPr.remove(e)
            tcw = OxmlElement('w:tcW'); tcw.set(qn('w:w'), str(sum(tw[ci:ci + n]))); tcw.set(qn('w:type'), 'dxa')
            tcPr.insert(0, tcw)
            if ri == 0:   # rule under header row
                bd = OxmlElement('w:tcBorders'); bt = OxmlElement('w:bottom')
                bt.set(qn('w:val'), 'single'); bt.set(qn('w:sz'), '4'); bt.set(qn('w:color'), '000000')
                bd.append(bt); tcPr.append(bd)
            ci += n
            for p in tc.findall(qn('w:p')):
                para = docx.text.paragraph.Paragraph(p, d._body)
                pf = para.paragraph_format
                pf.first_line_indent = Inches(0); pf.space_after = Pt(0); pf.space_before = Pt(0)
                pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in para.runs:
                    r.font.size = Pt(fsz)
                for r in list(p.iter(qn('w:r'))) + list(p.iter('{%s}r' % M_NS)):       # text, code and math runs
                    rp = r.find(qn('w:rPr'))
                    if rp is None:
                        rp = OxmlElement('w:rPr'); r.insert(0, rp)
                    for e in rp.findall(qn('w:sz')): rp.remove(e)
                    s = OxmlElement('w:sz'); s.set(qn('w:val'), str(2 * fsz)); rp.append(s)


# ------------------------------------------------------------------ walk body
children = list(body.iterchildren())
eq_no = 0
pending_table = None
wide_open = False
for el in children:
    if el.tag == qn('w:tbl'):
        if pending_table is not None:
            booktabs(docx.table.Table(el, d._body), *pending_table)
            pending_table = None
        continue
    if el.tag != qn('w:p'):
        continue
    txt = para_text(el)

    if txt.startswith('ABSTRACTMARK'):
        strip_marker(el, 'ABSTRACTMARK')
        para = runs_fmt(el, size=9, bold=True)
        para.paragraph_format.first_line_indent = Inches(0.17)
        r0 = OxmlElement('w:r'); rp = OxmlElement('w:rPr')
        for tag in ('w:b', 'w:i'): rp.append(OxmlElement(tag))
        s = OxmlElement('w:sz'); s.set(qn('w:val'), '18'); rp.append(s)
        r0.append(rp); tt = OxmlElement('w:t'); tt.text = 'Abstract—'; r0.append(tt)
        pPr = el.find(qn('w:pPr')); el.insert(1 if pPr is not None else 0, r0)
        continue
    if txt.startswith('KEYWORDSMARK'):
        strip_marker(el, 'KEYWORDSMARK')
        para = runs_fmt(el, size=9, bold=True)
        para.paragraph_format.space_after = Pt(6)
        r0 = OxmlElement('w:r'); rp = OxmlElement('w:rPr')
        for tag in ('w:b', 'w:i'): rp.append(OxmlElement(tag))
        s = OxmlElement('w:sz'); s.set(qn('w:val'), '18'); rp.append(s)
        r0.append(rp); tt = OxmlElement('w:t'); tt.text = 'Index Terms—'; r0.append(tt)
        pPr = el.find(qn('w:pPr')); el.insert(1 if pPr is not None else 0, r0)
        continue
    if txt.startswith('FIGMARK'):
        _, fname, kind = txt.split()
        wide = kind == 'WIDE'
        width = 6.9 if wide else (3.45 if fname == 'alg.png' else 3.35)
        pic = picture_par(fname, width)
        if wide:   # close the 2-column section before the figure
            brk, _ = mkp(); sect_break(brk, 2); el.addprevious(brk)
        el.addprevious(pic)
        body.remove(el)
        wide_open = wide
        continue
    if txt.startswith('FIGCAPMARK'):
        strip_marker(el, 'FIGCAPMARK')
        para = runs_fmt(el, size=8)
        pf = para.paragraph_format
        pf.first_line_indent = Inches(0); pf.space_after = Pt(8); pf.alignment = J
        if wide_open:
            brk, _ = mkp(); sect_break(brk, 1); el.addnext(brk); wide_open = False
        continue
    if txt.startswith('TABMARK'):
        _, num, kind, widths = txt.split()
        wide = kind == 'WIDE'
        pending_table = ([] if widths == '-' else widths.split(','), wide)
        cap, cpara = mkp('TABLE ' + num, size=8, align=WD_ALIGN_PARAGRAPH.CENTER, before=6)
        cpara.paragraph_format.keep_with_next = True
        if wide:
            brk, _ = mkp(); sect_break(brk, 2); el.addprevious(brk)
        el.addprevious(cap)
        body.remove(el)
        wide_open = wide
        continue
    if txt.startswith('TABCAPMARK'):
        strip_marker(el, 'TABCAPMARK')
        para = runs_fmt(el, size=8, small_caps=True)
        pf = para.paragraph_format
        pf.first_line_indent = Inches(0); pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.space_after = Pt(3); pf.keep_with_next = True
        continue
    if txt.startswith('TABENDMARK'):
        if wide_open:
            for t in el.iter(qn('w:t')): t.text = ''
            sect_break(el, 1); wide_open = False
        else:
            for t in el.iter(qn('w:t')): t.text = ''
            pf = docx.text.paragraph.Paragraph(el, d._body).paragraph_format
            pf.space_after = Pt(4); pf.space_before = Pt(0)
            rp = el.find(qn('w:pPr'))
            sp = OxmlElement('w:spacing'); sp.set(qn('w:line'), '80'); sp.set(qn('w:lineRule'), 'exact'); rp.append(sp)
        continue
    if txt.startswith('REFHEADMARK'):
        for t in el.iter(qn('w:t')): t.text = 'References'
        docx.text.paragraph.Paragraph(el, d._body).style = get_style('Heading 1')
        continue
    if txt.startswith('REFMARK'):
        strip_marker(el, 'REFMARK')
        para = runs_fmt(el, size=8)
        pf = para.paragraph_format
        pf.left_indent = Inches(0.25); pf.first_line_indent = Inches(-0.25); pf.space_after = Pt(1)
        for r in el.iter(qn('w:r')):
            rp = r.find(qn('w:rPr'))
            if rp is None: rp = OxmlElement('w:rPr'); r.insert(0, rp)
            if rp.find(qn('w:sz')) is None:
                s = OxmlElement('w:sz'); s.set(qn('w:val'), '16'); rp.append(s)
        continue

    # display equations -> centred math with right-aligned number
    omp = el.find('.//{%s}oMathPara' % M_NS)
    if omp is not None:
        eq_no += 1
        om = omp.find('{%s}oMath' % M_NS)
        new, npara = mkp(before=3, after=3)
        pPr = new.find(qn('w:pPr'))
        tabs = OxmlElement('w:tabs')
        for val, pos in (('center', COL_TW // 2), ('right', COL_TW)):
            tb = OxmlElement('w:tab'); tb.set(qn('w:val'), val); tb.set(qn('w:pos'), str(pos)); tabs.append(tb)
        pPr.append(tabs)
        r = OxmlElement('w:r'); r.append(OxmlElement('w:tab')); new.append(r)
        new.append(copy.deepcopy(om))
        r = OxmlElement('w:r'); r.append(OxmlElement('w:tab'))
        t = OxmlElement('w:t'); t.text = '(%d)' % eq_no; r.append(t); new.append(r)
        el.addprevious(new); body.remove(el)
        continue

    # list paragraphs: no first-line indent
    if el.find('.//' + qn('w:numPr')) is not None:
        docx.text.paragraph.Paragraph(el, d._body).paragraph_format.first_line_indent = Inches(0)

# ------------------------------------------------------------------ title + authors (single column)
first = body[0]
title, tpara = mkp('Validated, Event-Driven SMS Location Reporting for a Low-Cost GPS–GSM Vehicle Tracking Node',
                   size=24, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
first.addprevious(title)
authors = [('Rithikkaa S J', 'rithusj17@gmail.com'), ('Rahul S G', 'sg_rahul@ch.amrita.edu')]
tbl = d.add_table(rows=1, cols=2)
tbl_el = tbl._tbl
for i, (name, mail) in enumerate(authors):
    cell = tbl.rows[0].cells[i]
    cell.width = Inches(3.6)
    lines = [(name, 11, False), ('Dept. of Electronics and Communication Engineering', 10, True),
             ('Amrita Vishwa Vidyapeetham', 10, True), ('Chennai, India', 10, False), (mail, 10, False)]
    cell.paragraphs[0].text = ''
    for j, (txtl, sz, it) in enumerate(lines):
        p = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Inches(0); p.paragraph_format.space_after = Pt(0)
        r = p.add_run(txtl); r.font.size = Pt(sz); r.italic = it; r.font.name = TNR
tblPr = tbl_el.tblPr
jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'center'); tblPr.append(jc)
title.addnext(tbl_el)
brk, _ = mkp(after=6); sect_break(brk, 1, first=True); tbl_el.addnext(brk)

# final section = two columns
sp = body.find(qn('w:sectPr'))
if sp is not None: body.remove(sp)
body.append(make_sectpr(2))
for s in body.iter(qn('w:sectPr')):
    pass
# only the very first section starts on a new page; all later ones are continuous
first_sect = True
for s in body.iter(qn('w:sectPr')):
    t = s.find(qn('w:type'))
    t.set(qn('w:val'), 'nextPage' if first_sect else 'continuous'); first_sect = False

num = d.part.numbering_part.element if d.part._rels else None
try:
    nume = d.part.numbering_part.element
    for lvl in nume.iter(qn('w:lvl')):
        if lvl.get(qn('w:ilvl')) == '0':
            ind = lvl.find('.//' + qn('w:ind'))
            if ind is not None:
                ind.set(qn('w:left'), '300'); ind.set(qn('w:hanging'), '220')
except Exception as e:
    print('numbering not adjusted', e)
d.core_properties.title = 'Validated, Event-Driven SMS Location Reporting for a Low-Cost GPS–GSM Vehicle Tracking Node'
d.core_properties.author = 'Rithikkaa S J; Rahul S G'
d.save(OUT)
print('saved', OUT, 'equations', eq_no)
